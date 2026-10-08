"""Le dépôt Finance — balances mensuelles et taux de change, en classeurs Excel.

┌─ CE QUE CE JEU DOIT RENDRE TESTABLE ────────────────────────────────────────┐
│ La Finance dépose chaque mois, dans un dossier SharePoint, la balance de    │
│ chaque entité et le fichier des taux. Le consommateur (insights360, source  │
│ `depot_finance`) les lit par Graph avec la MÊME application Entra que les   │
│ groupes. Ce module fabrique ces classeurs AU DÉMARRAGE, à partir de code :  │
│ aucun binaire n'est commité, et chaque montant se relit ici.                │
│                                                                             │
│ Trois propriétés tiennent PAR CONSTRUCTION, et les tests les éprouvent :    │
│   1. chaque balance est ÉQUILIBRÉE au centime — elle est l'agrégat d'un     │
│      journal en partie double, pas une liste de soldes tirés au hasard ;    │
│   2. aucun compte n'est le préfixe strict d'un autre dans un même fichier : │
│      que des comptes FEUILLES, de longueur fixe dans chaque plan ;          │
│   3. deux constructions rendent les MÊMES octets — horodatages internes du  │
│      classeur et de l'archive fixés (cf. `_normaliser`).                    │
└─────────────────────────────────────────────────────────────────────────────┘

Les variantes INVALIDES (déséquilibrée, formule, sous-total, mauvaise devise…)
ne sont PAS dans le dossier par défaut : elles rendraient rouge la CI du
consommateur. Elles vivent dans `FIXTURES`, servies sous `/__fixtures`, et un
test les dépose explicitement par le plan de contrôle.
"""

from __future__ import annotations

import functools
import io
import random
import re
import zipfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import ROUND_HALF_UP, Decimal
from typing import Any

from openpyxl import Workbook

HOTE = "boreal-conseil.sharepoint.com"
CHEMIN_SITE = "/sites/depot-finance"
TITRE_SITE = "Dépôt Finance"
BIBLIOTHEQUE = "Documents"
DOSSIER = "Insights360"
SOUS_DOSSIER = "modeles"

#: Les en-têtes EXACTS — le consommateur refuse tout fichier qui s'en écarte,
#: et la variante `_entete` existe pour le vérifier.
ENTETE_BALANCE = ("entite", "mois", "compte", "libelle_compte", "debit", "credit", "devise")
ENTETE_TAUX = ("type_taux", "periode", "devise", "devise_pour_1_eur")

MOIS = tuple(f"2026-{m:02d}" for m in range(1, 7))

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_DOMAINE = "boreal-conseil.example"


# ── Les entités et leurs plans de comptes ────────────────────────────────────
#
# Le journal s'écrit UNE fois, par RÔLE (client, banque, chiffre d'affaires…) ;
# chaque entité décline ces rôles dans son propre référentiel. Trois plans,
# trois longueurs, trois langues — c'est ce que le consommateur devra mapper :
#   • NTE : plan comptable général français, comptes à six chiffres ;
#   • BOG : PUC colombien (Plan Único de Cuentas), sous-comptes à six chiffres ;
#   • MTL : plan « à la QuickBooks », comptes à quatre chiffres.
# À longueur fixe dans un plan, aucun compte ne peut être le préfixe d'un autre.


@dataclass(frozen=True)
class Parametres:
    """Les ordres de grandeur d'un mois, en UNITÉS de la devise de l'entité."""

    chiffre_affaires: int
    salaires: int
    taux_charges: float  # part patronale
    taux_retenues: float  # part salariale, retenue et reversée aux organismes
    loyer: int
    deplacements: int
    frais_bancaires: int
    taxe_ventes: float
    taxe_achats: float
    taxe_deplacements: float


@dataclass(frozen=True)
class Entite:
    code: str
    nom: str
    devise: str
    comptes: dict[str, tuple[str, str]]
    parametres: Parametres


ENTITES: dict[str, Entite] = {
    "NTE": Entite(
        code="NTE",
        nom="Boréal Conseil Nantes",
        devise="EUR",
        comptes={
            "fournisseur": ("401000", "Fournisseurs"),
            "client": ("411000", "Clients"),
            "personnel": ("421000", "Personnel - rémunérations dues"),
            "organismes": ("431000", "Sécurité sociale"),
            "taxe_deductible": ("445660", "TVA déductible sur autres biens et services"),
            "taxe": ("445710", "TVA collectée"),
            "banque": ("512000", "Banque"),
            "charge_interco": ("604800", "Achats de prestations intragroupe"),
            "loyer": ("613200", "Locations immobilières"),
            "deplacements": ("625100", "Voyages et déplacements"),
            "frais_bancaires": ("627800", "Frais et commissions bancaires"),
            "salaires": ("641100", "Salaires, appointements"),
            "charges_sociales": ("645100", "Cotisations à l'URSSAF"),
            "ca": ("706100", "Prestations de services"),
            "ca_interco": ("706800", "Prestations de services intragroupe"),
        },
        parametres=Parametres(
            chiffre_affaires=150_000,
            salaires=68_000,
            taux_charges=0.42,
            taux_retenues=0.22,
            loyer=9_500,
            deplacements=4_000,
            frais_bancaires=140,
            taxe_ventes=0.20,
            taxe_achats=0.20,
            taxe_deplacements=0.10,
        ),
    ),
    "BOG": Entite(
        code="BOG",
        nom="Boreal Conseil Bogota S.A.S",
        devise="COP",
        comptes={
            "caisse": ("110505", "Caja general"),
            "banque": ("111005", "Bancos - moneda nacional"),
            "client": ("130505", "Clientes nacionales"),
            "fournisseur": ("220505", "Proveedores nacionales"),
            "organismes": ("237005", "Aportes de nómina por pagar"),
            # IVA generado ET descontable sur le même sous-compte.
            "taxe": ("240805", "Impuesto sobre las ventas por pagar (IVA)"),
            "taxe_deductible": ("240805", "Impuesto sobre las ventas por pagar (IVA)"),
            "personnel": ("250505", "Salarios por pagar"),
            "ca": ("413505", "Ingresos por servicios de consultoría"),
            "ca_interco": ("413595", "Ingresos por servicios - intercompañía"),
            "salaires": ("510506", "Sueldos"),
            "charges_sociales": ("510568", "Aportes parafiscales y seguridad social"),
            "charge_interco": ("511095", "Honorarios - intercompañía"),
            "loyer": ("512010", "Arrendamientos - construcciones y edificaciones"),
            "deplacements": ("515505", "Gastos de viaje - alojamiento y manutención"),
            # Classe 53 : charges FINANCIÈRES (non opérationnelles) dans le PUC.
            "frais_bancaires": ("530505", "Gastos bancarios"),
        },
        parametres=Parametres(
            chiffre_affaires=400_000_000,
            salaires=180_000_000,
            taux_charges=0.21,
            taux_retenues=0.08,
            loyer=18_500_000,
            deplacements=2_600_000,
            frais_bancaires=650_000,
            taxe_ventes=0.19,
            taxe_achats=0.19,
            taxe_deplacements=0.0,
        ),
    ),
    "MTL": Entite(
        code="MTL",
        nom="Boreal Conseil Montreal",
        devise="CAD",
        comptes={
            "banque": ("1000", "Chequing"),
            "client": ("1200", "Accounts Receivable"),
            "fournisseur": ("2000", "Accounts Payable"),
            # Le net à payer ET les retenues à la source : QuickBooks les
            # regroupe dans un seul passif de paie.
            "personnel": ("2100", "Payroll Liabilities"),
            "organismes": ("2100", "Payroll Liabilities"),
            # Taxes perçues ET crédits de taxe sur intrants : un seul compte.
            "taxe": ("2200", "GST/QST Payable"),
            "taxe_deductible": ("2200", "GST/QST Payable"),
            "ca": ("4000", "Consulting Revenue"),
            "ca_interco": ("4090", "Intercompany Revenue"),
            "salaires": ("6000", "Salaries and Wages"),
            "charges_sociales": ("6010", "Payroll Taxes"),
            "charge_interco": ("6090", "Intercompany Management Fees"),
            "loyer": ("6200", "Rent"),
            "deplacements": ("6300", "Travel"),
            "frais_bancaires": ("7000", "Bank Fees"),
        },
        parametres=Parametres(
            chiffre_affaires=120_000,
            salaires=72_000,
            taux_charges=0.12,
            taux_retenues=0.25,
            loyer=8_200,
            deplacements=3_000,
            frais_bancaires=65,
            # TPS 5 % + TVQ 9,975 %.
            taxe_ventes=0.14975,
            taxe_achats=0.14975,
            taxe_deplacements=0.0,
        ),
    ),
}

#: Qui dépose quoi — l'identité que Graph rend dans `createdBy`/`lastModifiedBy`.
#: C'est de la DONNÉE PERSONNELLE (nom, courriel) : un consommateur qui oublie
#: `$select` la collecte, et c'est voulu qu'il la trouve ici.
DEPOSITAIRES: dict[str, tuple[str, str]] = {
    "NTE": ("Claire Rousseau", f"claire.rousseau@{_DOMAINE}"),
    "BOG": ("Camila Restrepo", f"camila.restrepo@{_DOMAINE}"),
    "MTL": ("Émilie Tremblay", f"emilie.tremblay@{_DOMAINE}"),
}


# ── Les taux ─────────────────────────────────────────────────────────────────
#
# Unités de devise pour UN euro. Pas de ligne EUR : l'euro est la devise de
# présentation, son taux est 1 par définition — et la variante `taux_2026_eur`
# existe pour vérifier qu'un consommateur refuse qu'on le lui donne.
TAUX_MOYENS: dict[str, dict[str, str]] = {
    "2026-01": {"CAD": "1.4712", "COP": "4528.41"},
    "2026-02": {"CAD": "1.4835", "COP": "4561.73"},
    "2026-03": {"CAD": "1.4968", "COP": "4497.26"},
    "2026-04": {"CAD": "1.5104", "COP": "4602.18"},
    "2026-05": {"CAD": "1.4887", "COP": "4655.94"},
    "2026-06": {"CAD": "1.4759", "COP": "4618.35"},
}
TAUX_BUDGET: dict[str, str] = {"CAD": "1.4900", "COP": "4550.00"}


def _taux(devise: str, mois: str) -> Decimal:
    """Le taux moyen du mois — le budgétaire pour un mois que le fichier ne
    couvre pas encore (les variantes de juillet)."""
    if devise == "EUR":
        return Decimal(1)
    return Decimal(TAUX_MOYENS.get(mois, TAUX_BUDGET)[devise])


# ── Le journal ───────────────────────────────────────────────────────────────

#: Une écriture : des lignes (rôle, débit, crédit) en CENTIMES. Elle est
#: équilibrée par construction ; `balance()` le vérifie quand même.
Ecriture = Sequence[tuple[str, int, int]]


def _flux_intragroupe(mois: str) -> dict[tuple[str, str], int]:
    """Les prestations intragroupe d'un mois, en centimes d'EURO.

    Tirées UNE fois pour le mois, puis converties dans la devise de chaque
    partie au taux moyen : les produits de l'une et les charges de l'autre se
    répondent à l'arrondi de change près — c'est ce qu'une élimination
    intragroupe doit retrouver.

    Clé : (vendeur, acheteur). Bogota et Montréal produisent pour Nantes ;
    Nantes leur refacture la direction.
    """
    alea = random.Random(f"depot-finance/intragroupe/{mois}")

    def tirer(base: int, ecart: float) -> int:
        return round(base * 100 * (1 + alea.uniform(-ecart, ecart)))

    return {
        ("BOG", "NTE"): tirer(24_000, 0.10),
        ("MTL", "NTE"): tirer(15_000, 0.10),
        ("NTE", "BOG"): tirer(5_000, 0.05),
        ("NTE", "MTL"): tirer(7_000, 0.05),
    }


def _convertir(centimes_eur: int, devise: str, mois: str) -> int:
    montant = Decimal(centimes_eur) * _taux(devise, mois)
    return int(montant.quantize(Decimal(1), rounding=ROUND_HALF_UP))


def _intragroupe(entite: Entite, mois: str) -> tuple[int, int]:
    """(produit, charge) intragroupe de l'entité, en centimes de SA devise."""
    flux = _flux_intragroupe(mois)
    produit = sum(m for (vendeur, _), m in flux.items() if vendeur == entite.code)
    charge = sum(m for (_, acheteur), m in flux.items() if acheteur == entite.code)
    return _convertir(produit, entite.devise, mois), _convertir(charge, entite.devise, mois)


def _journal(entite: Entite, mois: str) -> list[Ecriture]:
    """Le journal du mois — facturation, encaissements, paie, achats, banque.

    Déterministe : la graine est l'entité et le mois, rien d'autre. Les
    montants sont des MOUVEMENTS du mois, pas des soldes : un compte de bilan
    porte à la fois un débit et un crédit (le client est facturé PUIS encaissé).
    """
    p = entite.parametres
    alea = random.Random(f"depot-finance/{entite.code}/{mois}")

    def tirer(base: int, ecart: float) -> int:
        return round(base * 100 * (1 + alea.uniform(-ecart, ecart)))

    ca = tirer(p.chiffre_affaires, 0.08)
    taxe_ca = round(ca * p.taxe_ventes)
    # L'intragroupe est exonéré (export de services, autoliquidation) : pas de
    # taxe sur cette part de la facturation.
    produit_interco, charge_interco = _intragroupe(entite, mois)
    facture = ca + taxe_ca + produit_interco
    encaisse = round(facture * alea.uniform(0.82, 1.04))

    salaires = tirer(p.salaires, 0.03)
    charges = round(salaires * p.taux_charges)
    retenues = round(salaires * p.taux_retenues)
    net = salaires - retenues
    reverse = round((retenues + charges) * alea.uniform(0.97, 1.03))

    loyer = p.loyer * 100
    taxe_loyer = round(loyer * p.taxe_achats)
    deplacements = tirer(p.deplacements, 0.35)
    taxe_deplacements = round(deplacements * p.taxe_deplacements)
    frais = tirer(p.frais_bancaires, 0.30)

    ecritures: list[Ecriture] = [
        # Facturation du mois.
        [
            ("client", facture, 0),
            ("ca", 0, ca),
            ("taxe", 0, taxe_ca),
            ("ca_interco", 0, produit_interco),
        ],
        [("banque", encaisse, 0), ("client", 0, encaisse)],
        # Paie : brut et charges patronales en charge, net et retenues au passif.
        [
            ("salaires", salaires, 0),
            ("charges_sociales", charges, 0),
            ("personnel", 0, net),
            ("organismes", 0, retenues + charges),
        ],
        [("personnel", net, 0), ("banque", 0, net)],
        [("organismes", reverse, 0), ("banque", 0, reverse)],
        [
            ("loyer", loyer, 0),
            ("taxe_deductible", taxe_loyer, 0),
            ("fournisseur", 0, loyer + taxe_loyer),
        ],
        [("charge_interco", charge_interco, 0), ("fournisseur", 0, charge_interco)],
        [("frais_bancaires", frais, 0), ("banque", 0, frais)],
    ]
    achats = loyer + taxe_loyer + charge_interco
    if "caisse" in entite.comptes:
        # Bogota règle ses déplacements en espèces : la caisse est alimentée
        # depuis la banque, à la centaine de milliers de pesos supérieure.
        alimentation = -(-deplacements // 10_000_000) * 10_000_000
        ecritures += [
            [("caisse", alimentation, 0), ("banque", 0, alimentation)],
            [("deplacements", deplacements, 0), ("caisse", 0, deplacements)],
        ]
    else:
        ecritures.append(
            [
                ("deplacements", deplacements, 0),
                ("taxe_deductible", taxe_deplacements, 0),
                ("fournisseur", 0, deplacements + taxe_deplacements),
            ]
        )
        achats += deplacements + taxe_deplacements
    regle = round(achats * alea.uniform(0.85, 1.05))
    ecritures.append([("fournisseur", regle, 0), ("banque", 0, regle)])
    return ecritures


@dataclass(frozen=True)
class Ligne:
    """Une ligne de balance — montants en CENTIMES, jamais en flottants."""

    compte: str
    libelle: str
    debit: int
    credit: int


def balance(code: str, mois: str, *, ajustements: Sequence[Ecriture] = ()) -> list[Ligne]:
    """La balance des MOUVEMENTS du mois : l'agrégat du journal, par compte.

    Σ débit = Σ crédit au centime, parce que chaque écriture l'est. Le contrôle
    ci-dessous n'est pas décoratif : une écriture déséquilibrée par une faute
    de frappe dans `_journal` produirait un jeu de données qui ment sur ce
    qu'il prétend démontrer.
    """
    entite = ENTITES[code]
    totaux: dict[tuple[str, str], list[int]] = {}
    for ecriture in [*_journal(entite, mois), *ajustements]:
        if sum(d for _, d, _ in ecriture) != sum(c for _, _, c in ecriture):
            raise ValueError(f"écriture déséquilibrée pour {code} {mois} : {ecriture}")
        for role, debit, credit in ecriture:
            cumul = totaux.setdefault(entite.comptes[role], [0, 0])
            cumul[0] += debit
            cumul[1] += credit
    return sorted(
        (Ligne(compte, libelle, d, c) for (compte, libelle), (d, c) in totaux.items() if d or c),
        key=lambda ligne: ligne.compte,
    )


# ── Les classeurs ────────────────────────────────────────────────────────────

_FORMAT_MONTANT = "#,##0.00"
_FORMAT_TAUX = "0.0000"
_AUTEUR_CLASSEUR = "Boréal Conseil - Finance"


@dataclass(frozen=True)
class Feuille:
    titre: str
    lignes: list[list[Any]]
    #: Format numérique par colonne (1 = A), appliqué sous l'en-tête.
    formats: dict[int, str] = field(default_factory=dict)


#: L'horodatage de `docProps/core.xml` que `save()` écrase avec l'heure murale.
_MODIFIE = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def _normaliser(brut: bytes, horodatage: datetime) -> bytes:
    """Rend l'archive STABLE à l'octet près.

    ┌─ DEUX HORLOGES MURALES CACHÉES DANS UN .xlsx ─────────────────────────┐
    │ 1. openpyxl réécrit `dcterms:modified` avec l'heure courante à chaque │
    │    `save()`, quoi qu'on ait mis dans `workbook.properties` ;          │
    │ 2. `zipfile` date chaque membre de l'archive à l'instant de           │
    │    l'écriture.                                                        │
    │ Deux constructions du même classeur différaient donc d'une seconde à  │
    │ l'autre — et avec elles le `quickXorHash`, le `size`… et l'instantané │
    │ d'idempotence du consommateur, qui aurait vu « changer » un fichier   │
    │ que personne n'a touché. On réécrit l'archive avec un horodatage FIXE │
    │ et un système d'origine fixe (le défaut varie selon la plateforme).   │
    └───────────────────────────────────────────────────────────────────────┘
    """
    iso = horodatage.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    sortie = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(brut)) as entree,
        zipfile.ZipFile(sortie, "w", zipfile.ZIP_DEFLATED) as archive,
    ):
        for info in entree.infolist():
            donnees = entree.read(info)
            if info.filename == "docProps/core.xml":
                donnees = _MODIFIE.sub(rb"\g<1>" + iso + rb"\g<2>", donnees)
            membre = zipfile.ZipInfo(info.filename, date_time=horodatage.timetuple()[:6])
            membre.compress_type = zipfile.ZIP_DEFLATED
            membre.create_system = 0
            membre.external_attr = 0
            archive.writestr(membre, donnees)
    return sortie.getvalue()


def classeur(feuilles: Sequence[Feuille], horodatage: datetime) -> bytes:
    """Un classeur Excel, octet pour octet reproductible."""
    wb = Workbook()
    premiere = True
    for feuille in feuilles:
        ws = wb.active if premiere else wb.create_sheet()
        premiere = False
        assert ws is not None  # un Workbook neuf a toujours une feuille active
        ws.title = feuille.titre
        for ligne in feuille.lignes:
            ws.append(ligne)
        for colonne, format_ in feuille.formats.items():
            for rangee in ws.iter_rows(min_row=2, min_col=colonne, max_col=colonne):
                for cellule in rangee:
                    cellule.number_format = format_
    naif = horodatage.astimezone(UTC).replace(tzinfo=None)
    wb.properties.creator = _AUTEUR_CLASSEUR
    wb.properties.lastModifiedBy = _AUTEUR_CLASSEUR
    wb.properties.created = naif
    wb.properties.modified = naif
    tampon = io.BytesIO()
    wb.save(tampon)
    return _normaliser(tampon.getvalue(), naif)


def _rangees_balance(code: str, mois: str, devise: str, lignes: Sequence[Ligne]) -> list[list[Any]]:
    """Les rangées de la feuille `balance`. Le compte et le mois sont du TEXTE :
    `411000` saisi comme nombre perdrait ses zéros dans d'autres plans, et
    `2026-01` deviendrait une date."""
    return [
        [code, mois, ligne.compte, ligne.libelle, ligne.debit / 100, ligne.credit / 100, devise]
        for ligne in lignes
    ]


def classeur_balance(
    code: str,
    mois: str,
    horodatage: datetime,
    *,
    lignes: Sequence[Ligne] | None = None,
    entete: Sequence[str] = ENTETE_BALANCE,
    code_affiche: str | None = None,
    devise: str | None = None,
) -> bytes:
    entite = ENTITES[code]
    rangees = _rangees_balance(
        code_affiche or code,
        mois,
        devise or entite.devise,
        balance(code, mois) if lignes is None else lignes,
    )
    return classeur(
        [Feuille("balance", [list(entete), *rangees], {5: _FORMAT_MONTANT, 6: _FORMAT_MONTANT})],
        horodatage,
    )


def _rangees_taux() -> list[list[Any]]:
    rangees: list[list[Any]] = [
        ["moyen", mois, devise, float(taux)]
        for mois, taux_du_mois in TAUX_MOYENS.items()
        for devise, taux in sorted(taux_du_mois.items())
    ]
    rangees += [["budget", "2026", d, float(t)] for d, t in sorted(TAUX_BUDGET.items())]
    return rangees


def classeur_taux(horodatage: datetime, *, supplementaires: Sequence[list[Any]] = ()) -> bytes:
    """`taux_2026.xlsx` — la période est du TEXTE, y compris l'année budgétaire."""
    return classeur(
        [
            Feuille(
                "taux", [list(ENTETE_TAUX), *_rangees_taux(), *supplementaires], {4: _FORMAT_TAUX}
            )
        ],
        horodatage,
    )


def _modele(horodatage: datetime) -> bytes:
    """Le modèle que la Finance recopie chaque mois — un en-tête et une notice.
    Il vit dans un SOUS-DOSSIER, que le consommateur doit ignorer."""
    return classeur(
        [
            Feuille("balance", [list(ENTETE_BALANCE)]),
            Feuille(
                "lisez-moi",
                [
                    ["Une feuille `balance`, un fichier par entité et par mois."],
                    ["Nom : balance_<CODE>_<AAAA-MM>.xlsx — CODE parmi NTE, BOG, MTL."],
                    ["Montants : mouvements du mois, au centime, sans formule."],
                ],
            ),
        ],
        horodatage,
    )


# ── Le jeu par défaut ────────────────────────────────────────────────────────


def _instant(texte: str) -> datetime:
    return datetime.strptime(texte, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=UTC)


def iso(instant: datetime) -> str:
    return instant.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")


#: L'heure de dépôt de chaque entité, le 6 du mois suivant — chacune dans sa
#: matinée locale. FIXE : un horodatage tiré de l'horloge murale rendrait
#: l'instantané d'idempotence du consommateur différent à chaque démarrage.
_HEURE_DE_DEPOT = {"NTE": "08:14", "BOG": "14:37", "MTL": "13:52"}


def depot_de_la_balance(code: str, mois: str) -> datetime:
    annee, numero = (int(x) for x in mois.split("-"))
    annee, numero = (annee + 1, 1) if numero == 12 else (annee, numero + 1)
    return _instant(f"{annee}-{numero:02d}-06T{_HEURE_DE_DEPOT[code]}:00Z")


@dataclass(frozen=True)
class Fichier:
    chemin: str
    contenu: bytes
    cree: str
    modifie: str
    #: Code d'entité du déposant (cf. DEPOSITAIRES).
    depositaire: str
    #: Le `n` de `eTag`/`cTag` : le nombre de versions du contenu.
    version: int = 1


#: Les dossiers, du plus haut au plus profond, avec leur date de création.
DOSSIERS: tuple[tuple[str, str], ...] = (
    (DOSSIER, "2025-12-15T09:55:00Z"),
    (f"{DOSSIER}/{SOUS_DOSSIER}", "2025-12-15T10:00:00Z"),
)


def construire_jeu_par_defaut() -> list[Fichier]:
    """Les 18 balances, le fichier des taux et le modèle — SANS cache.

    `jeu_par_defaut()` est la version mise en cache que sert le mock ; celle-ci
    existe pour qu'un test puisse construire deux fois et comparer les octets.
    """
    fichiers = [
        Fichier(
            chemin=f"{DOSSIER}/balance_{code}_{mois}.xlsx",
            contenu=classeur_balance(code, mois, depot_de_la_balance(code, mois)),
            cree=iso(depot_de_la_balance(code, mois)),
            modifie=iso(depot_de_la_balance(code, mois)),
            depositaire=code,
        )
        for code in ENTITES
        for mois in MOIS
    ]
    # Le fichier des taux est créé en janvier et COMPLÉTÉ chaque mois : six
    # versions, d'où un `cTag` en `,6` — un `n` qui n'est pas toujours 1.
    taux_modifie = _instant("2026-07-03T09:30:00Z")
    fichiers.append(
        Fichier(
            chemin=f"{DOSSIER}/taux_2026.xlsx",
            contenu=classeur_taux(taux_modifie),
            cree="2026-01-05T09:00:00Z",
            modifie=iso(taux_modifie),
            depositaire="NTE",
            version=6,
        )
    )
    modele = _instant("2025-12-15T10:05:00Z")
    fichiers.append(
        Fichier(
            chemin=f"{DOSSIER}/{SOUS_DOSSIER}/modele_balance.xlsx",
            contenu=_modele(modele),
            cree=iso(modele),
            modifie=iso(modele),
            depositaire="NTE",
        )
    )
    return fichiers


@functools.cache
def jeu_par_defaut() -> tuple[Fichier, ...]:
    return tuple(construire_jeu_par_defaut())


# ── Les variantes : invalides, et deux valides de plus ───────────────────────
#
# UN défaut par variante, et un seul : un test du consommateur qui attend
# « refusé pour déséquilibre » ne doit pas passer parce que le fichier était
# AUSSI refusé pour son en-tête. Chaque variante est donc valide en tout
# point, sauf celui qu'elle nomme.

_JUILLET = _instant("2026-08-06T08:14:00Z")


def _fichier_verrou(proprietaire: str) -> bytes:
    """Le fichier « propriétaire » qu'Excel pose à côté d'un classeur ouvert.

    165 octets : la longueur du nom, le nom en ASCII complété d'espaces, puis
    le même en UTF-16. Ce n'est PAS un classeur : l'ouvrir comme tel lève une
    erreur de zip.
    """
    n = len(proprietaire)
    contenu = bytes([n]) + proprietaire.encode("ascii", "replace").ljust(53, b" ")
    contenu += bytes([n, 0]) + proprietaire.encode("utf-16-le")
    contenu += b" \x00" * ((165 - len(contenu)) // 2)
    return contenu.ljust(165, b"\x00")[:165]


def _avec_debit_augmente(code: str, mois: str, compte: str, centimes: int) -> list[Ligne]:
    return [
        Ligne(lg.compte, lg.libelle, lg.debit + centimes, lg.credit) if lg.compte == compte else lg
        for lg in balance(code, mois)
    ]


def _desequilibree() -> bytes:
    # +1 000,00 COP au débit des frais de voyage, sans contrepartie.
    lignes = _avec_debit_augmente("BOG", "2026-07", "515505", 100_000)
    return classeur_balance("BOG", "2026-07", _JUILLET, lignes=lignes)


def _formule() -> bytes:
    """Le débit des déplacements saisi comme une SOMME (`=a+b`).

    Sa valeur, une fois calculée, est la bonne : la balance tomberait juste.
    Mais openpyxl n'écrit aucune valeur en cache pour une formule — un lecteur
    en `data_only=True` lit `None`, un lecteur brut lit la chaîne `=…`. Les
    deux doivent refuser le fichier, pas l'un des deux.
    """
    lignes = balance("MTL", "2026-07")
    rangees = _rangees_balance("MTL", "2026-07", "CAD", lignes)
    for rangee, ligne in zip(rangees, lignes, strict=True):
        if ligne.compte == "6300":
            a, b = ligne.debit * 3 // 5, ligne.debit - ligne.debit * 3 // 5
            rangee[4] = f"={a // 100}.{a % 100:02d}+{b // 100}.{b % 100:02d}"
    return classeur(
        [Feuille("balance", [list(ENTETE_BALANCE), *rangees], {5: _FORMAT_MONTANT})], _JUILLET
    )


def _sous_total() -> bytes:
    """`6411` ET `641100` dans le même fichier — un compte parent mouvementé à
    côté de son enfant. Équilibrée : la paire préfixe est son SEUL défaut."""
    lignes: list[Ligne] = []
    for lg in balance("NTE", "2026-07"):
        if lg.compte == "641100":
            part = lg.debit * 12 // 100
            lignes += [
                Ligne("6411", "Salaires, appointements (sous-total)", part, 0),
                Ligne(lg.compte, lg.libelle, lg.debit - part, lg.credit),
            ]
        else:
            lignes.append(lg)
    return classeur_balance("NTE", "2026-07", _JUILLET, lignes=lignes)


def _corrigee() -> bytes:
    # Une régularisation de loyer de 250,00 € : même mois, autres octets.
    ajustement = [("loyer", 25_000, 0), ("fournisseur", 0, 25_000)]
    lignes = balance("NTE", "2026-06", ajustements=[ajustement])
    return classeur_balance("NTE", "2026-06", _instant("2026-07-09T16:02:00Z"), lignes=lignes)


def _notes() -> bytes:
    return (
        "date;auteur;note\n"
        "2026-07-02;Claire Rousseau;Balance BOG de juin re-déposée après correction\n"
    ).encode()


@dataclass(frozen=True)
class Fixture:
    nom: str
    #: Ce qui cloche — la règle que le consommateur doit appliquer pour la refuser.
    defaut: str
    fabriquer: Callable[[], bytes]
    #: Une variante VALIDE (un mois de plus, une correction) : à accepter.
    valide: bool = False


#: LA liste des variantes — le README la reprend, `/__fixtures` la sert.
FIXTURES: tuple[Fixture, ...] = (
    Fixture(
        "balance_BOG_2026-07_desequilibree.xlsx",
        "Σ débit - Σ crédit = 1 000,00 COP (compte 515505)",
        _desequilibree,
    ),
    Fixture(
        "balance_MTL_2026-07_formule.xlsx",
        "le débit du compte 6300 est une formule (=a+b), sans valeur en cache",
        _formule,
    ),
    Fixture(
        "balance_NTE_2026-07_sous_total.xlsx",
        "6411 et 641100 dans le même fichier : un compte préfixe d'un autre",
        _sous_total,
    ),
    Fixture(
        "balance_BOG_2026-07_mauvaise_devise.xlsx",
        "devise USD sur toutes les lignes — BOG tient ses comptes en COP",
        lambda: classeur_balance("BOG", "2026-07", _JUILLET, devise="USD"),
    ),
    Fixture(
        "balance_NTE_2026-07_entete.xlsx",
        "en-tête saisi à la main (accents, majuscules, libellés) au lieu de l'en-tête exact",
        lambda: classeur_balance(
            "NTE",
            "2026-07",
            _JUILLET,
            entete=("Entité", "Mois", "Compte", "Libellé", "Débit", "Crédit", "Devise"),
        ),
    ),
    Fixture(
        "balance_XXX_2026-07.xlsx",
        "code d'entité inconnu (XXX), dans le nom comme dans la colonne entite",
        lambda: classeur_balance("NTE", "2026-07", _JUILLET, code_affiche="XXX"),
    ),
    Fixture(
        "taux_2026_eur.xlsx",
        "une ligne EUR (moyen 2026-01, taux 1) — l'euro n'a pas de taux contre lui-même",
        lambda: classeur_taux(_JUILLET, supplementaires=[["moyen", "2026-01", "EUR", 1.0]]),
    ),
    Fixture(
        "~$balance_NTE_2026-01.xlsx",
        "fichier de verrou Excel (~$…) : 165 octets, pas un classeur",
        lambda: _fichier_verrou("Claire Rousseau"),
    ),
    Fixture(
        "notes.csv",
        "ni balance ni taux : un fichier hors convention de nommage",
        _notes,
    ),
    Fixture(
        "balance_NTE_2026-07.xlsx",
        "aucun — un mois de plus, à accepter (dépôt incrémental)",
        lambda: classeur_balance("NTE", "2026-07", _JUILLET),
        valide=True,
    ),
    Fixture(
        "balance_NTE_2026-06_corrigee.xlsx",
        "aucun — juin re-déposé corrigé (+250,00 € de loyer), à déposer SOUS le nom "
        "balance_NTE_2026-06.xlsx pour éprouver l'écrasement",
        _corrigee,
        valide=True,
    ),
)


@functools.cache
def fixture(nom: str) -> bytes | None:
    trouvee = next((f for f in FIXTURES if f.nom == nom), None)
    return None if trouvee is None else trouvee.fabriquer()
