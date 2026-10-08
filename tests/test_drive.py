"""La surface SharePoint — le dépôt Finance que lit la source `depot_finance`.

Chaque test vise un comportement du VRAI Graph qu'un mock gentil cacherait :
le 403 d'une application `Sites.Selected` sans droit sur le site, la
pagination de `/children`, le jeu de propriétés par défaut qui porte une
personne, l'empreinte `quickXorHash` seule, la redirection 302 vers une URL
pré-authentifiée à suivre SANS le Bearer. Puis le jeu de données lui-même :
chaque balance équilibrée au centime, sans compte préfixe d'un autre, et
reproductible à l'octet près.
"""

from __future__ import annotations

import base64
import hashlib
import re
import zipfile
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlsplit

import pytest
import yaml
from conftest import ADMIN, app_module, drive_module
from openpyxl import load_workbook

from entra_mock.dataset import depot_finance as jeu

SITE = "/v1.0/sites/boreal-conseil.sharepoint.com:/sites/depot-finance"
DRIVE = drive_module.DRIVE_ID
DOSSIER = f"/v1.0/drives/{DRIVE}/root:/Insights360:/children"
NOMS_PAR_DEFAUT = {
    *(f"balance_{code}_{mois}.xlsx" for code in ("NTE", "BOG", "MTL") for mois in jeu.MOIS),
    "taux_2026.xlsx",
    "modeles",
}


@pytest.fixture(autouse=True)
def lecteur_neuf():
    """Un lecteur REMIS AU JEU PAR DÉFAUT — avant et après, pour qu'un dépôt
    ou un droit retiré ne fuie pas d'un test à l'autre."""
    drive_module.depot.reinitialiser()
    yield
    drive_module.depot.reinitialiser()


def _relatif(url: str) -> str:
    return url.replace("http://testserver", "")


def tous_les_enfants(client, auth, url: str = DOSSIER) -> list[dict[str, Any]]:
    """Les enfants en SUIVANT `@odata.nextLink` jusqu'au bout."""
    elements: list[dict[str, Any]] = []
    while url:
        reponse = client.get(_relatif(url), headers=auth)
        assert reponse.status_code == 200, reponse.text
        corps = reponse.json()
        elements.extend(corps["value"])
        url = corps.get("@odata.nextLink", "")
    return elements


def element(client, auth, nom: str) -> dict[str, Any]:
    reponse = client.get(f"/v1.0/drives/{DRIVE}/root:/Insights360/{nom}", headers=auth)
    assert reponse.status_code == 200, reponse.text
    return dict(reponse.json())


def telecharger(client, auth, item_id: str) -> bytes:
    """Le geste CORRECT : `/content` sans suivre, puis `Location` SANS jeton."""
    redirection = client.get(
        f"/v1.0/drives/{DRIVE}/items/{item_id}/content", headers=auth, follow_redirects=False
    )
    assert redirection.status_code == 302, redirection.text
    reponse = client.get(redirection.headers["location"])
    assert reponse.status_code == 200, reponse.text
    return bytes(reponse.content)


def deposer(client, chemin: str, contenu: bytes):
    return client.put(f"/__admin/drive/files/{chemin}", content=contenu, headers=ADMIN)


# ── Le site et ses bibliothèques ─────────────────────────────────────────────


def test_le_site_par_son_chemin_rend_un_identifiant_composite(client, auth):
    r = client.get(SITE, headers=auth)
    assert r.status_code == 200, r.text
    site = r.json()
    hote, collection, web = site["id"].split(",")
    assert hote == "boreal-conseil.sharepoint.com"
    assert re.fullmatch(r"[0-9a-f-]{36}", collection) and re.fullmatch(r"[0-9a-f-]{36}", web)
    assert site["webUrl"] == "https://boreal-conseil.sharepoint.com/sites/depot-finance"
    assert {"displayName", "name"} <= set(site)
    # Le même site par son identifiant.
    assert client.get(f"/v1.0/sites/{site['id']}", headers=auth).json()["id"] == site["id"]


@pytest.mark.parametrize(
    "chemin",
    [
        "/v1.0/sites/boreal-conseil.sharepoint.com:/sites/depot-finance:",
        "/v1.0/sites/boreal-conseil.sharepoint.com:/sites/Depot-Finance",
        "/v1.0/sites/boreal-conseil.sharepoint.com%3A/sites/depot-finance",
    ],
)
def test_les_variantes_d_adressage_du_site(client, auth, chemin):
    """Deux-points final, casse, deux-points encodé : Graph accepte les trois."""
    assert client.get(chemin, headers=auth).status_code == 200


def test_site_inconnu_404_itemNotFound(client, auth):
    r = client.get("/v1.0/sites/boreal-conseil.sharepoint.com:/sites/autre", headers=auth)
    assert r.status_code == 404
    assert r.json()["error"]["code"] == "itemNotFound"


def test_hote_d_un_autre_locataire_400(client, auth):
    r = client.get("/v1.0/sites/contoso.sharepoint.com:/sites/depot-finance", headers=auth)
    assert r.status_code == 400
    assert r.json()["error"]["code"] == "invalidRequest"


def test_sans_droit_sur_le_site_403_accessDenied_PARTOUT(client, auth, monkeypatch):
    """L'application `Sites.Selected` à qui personne n'a accordé le site.

    Le jeton est VALIDE — c'est le site qui est fermé. Le 403 doit tomber sur
    le site comme sur tout ce qu'il contient, avec l'enveloppe de Graph : c'est
    ce qui permet au consommateur de dire « demander le droit » plutôt que
    « renouveler le jeton » (401) ou « corriger le chemin » (404).
    """
    monkeypatch.setattr(drive_module, "SITE_GRANT", "none")
    drive_module.depot.reinitialiser()
    for url in (
        SITE,
        f"/v1.0/sites/{drive_module.SITE_ID}/drive",
        DOSSIER,
        f"/v1.0/drives/{DRIVE}/root:/Insights360/taux_2026.xlsx:/content",
    ):
        r = client.get(url, headers=auth, follow_redirects=False)
        assert r.status_code == 403, url
        erreur = r.json()["error"]
        assert erreur["code"] == "accessDenied"
        assert erreur["message"] == "Access denied"
        assert {"date", "request-id", "client-request-id"} <= set(erreur["innerError"])


def test_le_droit_se_retire_et_se_rend_a_chaud(client, auth):
    """Un conteneur de CI ne se redémarre pas entre deux tests : le pendant
    runtime de ENTRA_MOCK_SITE_GRANT passe par le plan de contrôle."""
    assert client.post("/__admin/drive/grant", json={"grant": "none"}, headers=ADMIN).is_success
    assert client.get(SITE, headers=auth).status_code == 403
    assert client.post("/__admin/drive/grant", json={"grant": "read"}, headers=ADMIN).is_success
    assert client.get(SITE, headers=auth).status_code == 200


def test_la_bibliotheque_par_defaut_et_la_liste(client, auth):
    site_id = client.get(SITE, headers=auth).json()["id"]
    lecteur = client.get(f"/v1.0/sites/{site_id}/drive", headers=auth).json()
    assert lecteur["name"] == "Documents"
    assert lecteur["driveType"] == "documentLibrary"
    assert lecteur["id"].startswith("b!") and len(lecteur["id"]) == 66
    lecteurs = client.get(f"/v1.0/sites/{site_id}/drives", headers=auth).json()
    assert [d["id"] for d in lecteurs["value"]] == [lecteur["id"]]
    assert client.get(f"/v1.0/drives/{lecteur['id']}", headers=auth).json()["name"] == "Documents"


# ── Les enfants du dossier ───────────────────────────────────────────────────


def test_children_PAGINE_par_curseur_opaque(client, auth):
    """Deux par page : un client qui ignore `@odata.nextLink` ne voit que
    `balance_BOG_2026-01` et `-02` — sans la moindre erreur."""
    premiere = client.get(DOSSIER, headers=auth).json()
    assert len(premiere["value"]) == 2
    suivant = premiere["@odata.nextLink"]
    assert suivant.startswith("http://testserver/v1.0/drives/")
    jeton = parse_qs(urlsplit(suivant).query)["$skiptoken"][0]
    assert not jeton.isdigit(), "le curseur doit être OPAQUE, pas un rang"

    tous = tous_les_enfants(client, auth)
    noms = [e["name"] for e in tous]
    assert len(noms) == len(set(noms)) == 20
    assert set(noms) == NOMS_PAR_DEFAUT


def test_top_est_un_plafond_pas_une_promesse(client, auth):
    """`$top=500` ne court-circuite PAS la pagination : le serveur rend moins."""
    grande = client.get(f"{DOSSIER}?$top=500", headers=auth).json()
    assert len(grande["value"]) == 2
    assert "$top=500" in grande["@odata.nextLink"]
    assert len(client.get(f"{DOSSIER}?$top=1", headers=auth).json()["value"]) == 1
    assert client.get(f"{DOSSIER}?$top=zero", headers=auth).status_code == 400


def test_un_curseur_illisible_400(client, auth):
    assert client.get(f"{DOSSIER}?$skiptoken=pas-un-jeton", headers=auth).status_code == 400


def test_le_select_MINIMISE_et_suit_le_lien(client, auth):
    """`$select` rend EXACTEMENT les propriétés demandées — pas d'`id` offert,
    pas de personne — et il est recopié dans le lien suivant."""
    demande = "id,name,size,cTag,lastModifiedDateTime,file,folder"
    tous = tous_les_enfants(client, auth, f"{DOSSIER}?$select={demande}")
    assert len(tous) == 20
    for item in tous:
        assert set(item) <= set(demande.split(",")), item
        assert "createdBy" not in item and "lastModifiedBy" not in item
    # Sans `id` demandé, pas d'`id` rendu.
    seul = client.get(f"{DOSSIER}?$select=name", headers=auth).json()
    assert all(set(item) == {"name"} for item in seul["value"])


def test_un_select_inconnu_400(client, auth):
    r = client.get(f"{DOSSIER}?$select=name,sha256Hash", headers=auth)
    assert r.status_code == 400
    assert "sha256Hash" in r.json()["error"]["message"]


def test_sans_select_le_jeu_par_defaut_expose_une_PERSONNE(client, auth):
    """Le défaut que `$select` existe pour éviter : nom et courriel du déposant."""
    item = client.get(DOSSIER, headers=auth).json()["value"][0]
    for cle in ("createdBy", "lastModifiedBy"):
        personne = item[cle]["user"]
        assert personne["displayName"]
        assert personne["email"].endswith("@boreal-conseil.example")
    assert {"parentReference", "webUrl", "fileSystemInfo", "@microsoft.graph.downloadUrl"} <= set(
        item
    )


def test_SEUL_quickXorHash_jamais_sha(client, auth):
    """SharePoint ne calcule ni sha1 ni sha256 : un consommateur qui en lit un
    tombe sur une clé absente, ici comme en production."""
    for item in tous_les_enfants(client, auth):
        if item["name"] == "modeles":
            assert item["folder"]["childCount"] == 1
            assert "file" not in item
            continue
        assert set(item["file"]["hashes"]) == {"quickXorHash"}
        assert item["file"]["mimeType"] == jeu.MIME_XLSX
        assert "folder" not in item


def test_etag_et_ctag_ont_la_forme_de_sharepoint(client, auth):
    item = element(client, auth, "balance_NTE_2026-01.xlsx")
    assert re.fullmatch(r'"\{[0-9A-F-]{36}\},1"', item["eTag"])
    assert re.fullmatch(r'"c:\{[0-9A-F-]{36}\},1"', item["cTag"])
    assert item["eTag"][2:38] == item["cTag"][4:40]
    # Le fichier des taux a été complété six fois : `n` n'est pas toujours 1.
    assert element(client, auth, "taux_2026.xlsx")["cTag"].endswith(',6"')


def test_le_sous_dossier_est_un_dossier_que_l_on_peut_lister(client, auth):
    modeles = element(client, auth, "modeles")
    assert "folder" in modeles and "file" not in modeles
    enfants = tous_les_enfants(client, auth, f"/v1.0/drives/{DRIVE}/items/{modeles['id']}/children")
    assert [e["name"] for e in enfants] == ["modele_balance.xlsx"]


def test_element_par_identifiant_egal_element_par_chemin(client, auth):
    par_chemin = element(client, auth, "balance_MTL_2026-03.xlsx")
    par_id = client.get(f"/v1.0/drives/{DRIVE}/items/{par_chemin['id']}", headers=auth).json()
    for cle in ("id", "name", "size", "eTag", "cTag", "file", "lastModifiedDateTime"):
        assert par_id[cle] == par_chemin[cle]
    assert re.fullmatch(r"01[A-Z2-7]{32}", par_id["id"])


def test_la_casse_du_chemin_est_indifferente(client, auth):
    assert (
        client.get(DOSSIER.replace("Insights360", "insights360"), headers=auth).status_code == 200
    )


def test_dossier_inconnu_404_itemNotFound(client, auth):
    r = client.get(f"/v1.0/drives/{DRIVE}/root:/Finance:/children", headers=auth)
    assert r.status_code == 404
    assert r.json()["error"]["code"] == "itemNotFound"
    assert (
        client.get("/v1.0/drives/b!inconnu/root:/Insights360:/children", headers=auth).status_code
        == 404
    )
    assert client.get(f"/v1.0/drives/{DRIVE}/items/01INCONNU", headers=auth).status_code == 404


# ── Le contenu : 302, puis une URL pré-authentifiée ──────────────────────────


def test_content_rend_302_vers_une_url_pre_authentifiee(client, auth):
    item = element(client, auth, "balance_NTE_2026-01.xlsx")
    r = client.get(
        f"/v1.0/drives/{DRIVE}/items/{item['id']}/content", headers=auth, follow_redirects=False
    )
    assert r.status_code == 302
    lieu = urlsplit(r.headers["location"])
    assert lieu.scheme == "http" and lieu.netloc == "testserver", "Location doit être ABSOLUE"
    assert lieu.path == "/sites/depot-finance/_layouts/15/download.aspx"
    parametres = parse_qs(lieu.query)
    assert {"UniqueId", "tempauth"} <= set(parametres)

    # Suivie SANS jeton : les octets, avec le bon type.
    telechargement = client.get(r.headers["location"])
    assert telechargement.status_code == 200
    assert telechargement.headers["content-type"] == jeu.MIME_XLSX
    assert len(telechargement.content) == item["size"]


def test_content_par_chemin_aussi(client, auth):
    r = client.get(
        f"/v1.0/drives/{DRIVE}/root:/Insights360/taux_2026.xlsx:/content",
        headers=auth,
        follow_redirects=False,
    )
    assert r.status_code == 302


def test_renvoyer_le_Bearer_a_l_url_de_telechargement_401(client, auth):
    """LA faute : l'autorisation est DANS l'URL. Le jeton Graph n'a rien à y
    faire — un client qui le renvoie est refusé."""
    item = element(client, auth, "balance_NTE_2026-01.xlsx")
    r = client.get(
        f"/v1.0/drives/{DRIVE}/items/{item['id']}/content", headers=auth, follow_redirects=False
    )
    assert client.get(r.headers["location"], headers=auth).status_code == 401
    # Le piège en vrai : suivre la redirection AVEC le jeton (même origine ici).
    suivie = client.get(f"/v1.0/drives/{DRIVE}/items/{item['id']}/content", headers=auth)
    assert suivie.status_code == 401


def test_tempauth_falsifie_ou_expire_401(client, auth, monkeypatch):
    item = element(client, auth, "balance_NTE_2026-01.xlsx")
    url = item["@microsoft.graph.downloadUrl"]
    assert client.get(url).status_code == 200
    jeton = parse_qs(urlsplit(url).query)["tempauth"][0]
    falsifie = url.replace(jeton, jeton[:-4] + ("AAAA" if not jeton.endswith("AAAA") else "BBBB"))
    assert client.get(falsifie).status_code == 401
    assert client.get(url.replace(jeton, "inconnu")).status_code == 401
    # Un tempauth valide pour un AUTRE fichier ne l'ouvre pas.
    autre = element(client, auth, "balance_NTE_2026-02.xlsx")["@microsoft.graph.downloadUrl"]
    jeton_autre = parse_qs(urlsplit(autre).query)["tempauth"][0]
    assert client.get(url.replace(jeton, jeton_autre)).status_code == 401

    monkeypatch.setattr(drive_module, "TEMPAUTH_SECONDS", -1)
    expire = element(client, auth, "balance_NTE_2026-01.xlsx")["@microsoft.graph.downloadUrl"]
    assert client.get(expire).status_code == 401


def test_quick_xor_hash_vecteurs_de_reference():
    """Vecteurs relevés contre l'implémentation de référence (paquet PyPI
    `quickxorhash` 1.0.5, le 2026-10-08)."""
    vecteurs = {
        b"": "AAAAAAAAAAAAAAAAAAAAAAAAAAA=",
        b"a": "YQAAAAAAAAAAAAAAAQAAAAAAAAA=",
        b"hello world": "aCgDG9jwBhDc4Q1yawMZAAAAAAA=",
        bytes(range(256)) * 3: "rxAOGe1RimTF/e+k/m0O5nnSZT8=",
    }
    for donnees, attendu in vecteurs.items():
        assert drive_module.quick_xor_hash(donnees) == attendu


def test_le_quickXorHash_servi_est_celui_des_octets(client, auth):
    item = element(client, auth, "taux_2026.xlsx")
    octets = telecharger(client, auth, item["id"])
    assert item["file"]["hashes"]["quickXorHash"] == drive_module.quick_xor_hash(octets)
    assert len(base64.b64decode(item["file"]["hashes"]["quickXorHash"])) == 20


# ── Le plan de contrôle ──────────────────────────────────────────────────────


def test_le_plan_de_controle_exige_son_jeton(client):
    assert client.post("/__admin/drive/reset").status_code == 403
    assert (
        client.post("/__admin/drive/reset", headers={"X-Mock-Admin-Token": "faux"}).status_code
        == 403
    )
    assert client.post("/__admin/drive/reset", headers=ADMIN).status_code == 200


def test_ecraser_GARDE_l_id_et_avance_ctag_etag_et_l_horloge(client, auth):
    """La Finance re-dépose juin corrigé, sous le MÊME nom.

    Même élément SharePoint (même id), contenu neuf : `n` passe à 2 dans
    `eTag` et `cTag`, `lastModifiedDateTime` avance — sur l'horloge du mock,
    pas l'horloge murale, pour que l'instantané du consommateur reste stable.
    """
    avant = element(client, auth, "balance_NTE_2026-06.xlsx")
    corrigee = client.get("/__fixtures/depot_finance/balance_NTE_2026-06_corrigee.xlsx").content
    r = deposer(client, "Insights360/balance_NTE_2026-06.xlsx", corrigee)
    assert r.status_code == 200, r.text
    apres = element(client, auth, "balance_NTE_2026-06.xlsx")

    assert apres["id"] == avant["id"]
    assert apres["cTag"] == avant["cTag"].replace(",1", ",2")
    assert apres["eTag"] == avant["eTag"].replace(",1", ",2")
    assert apres["lastModifiedDateTime"] == "2026-07-15T08:01:00Z"
    assert apres["lastModifiedDateTime"] > avant["lastModifiedDateTime"]
    assert apres["createdDateTime"] == avant["createdDateTime"]
    assert apres["file"]["hashes"] != avant["file"]["hashes"]
    assert telecharger(client, auth, apres["id"]) == corrigee


def test_deposer_un_nouveau_fichier_201(client, auth):
    juillet = client.get("/__fixtures/depot_finance/balance_NTE_2026-07.xlsx").content
    r = deposer(client, "Insights360/balance_NTE_2026-07.xlsx", juillet)
    assert r.status_code == 201, r.text
    assert r.json()["cTag"].endswith(',1"')
    assert r.json()["lastModifiedBy"]["user"]["displayName"] == "Claire Rousseau"
    noms = {e["name"] for e in tous_les_enfants(client, auth)}
    assert noms == NOMS_PAR_DEFAUT | {"balance_NTE_2026-07.xlsx"}


def test_l_horloge_du_mock_est_deterministe(client):
    """La même suite d'écritures rend les mêmes horodatages, run après run."""

    def horodatages() -> list[str]:
        return [
            deposer(client, f"Insights360/f{i}.csv", b"x").json()["lastModifiedDateTime"]
            for i in range(3)
        ]

    premiers = horodatages()
    assert premiers == ["2026-07-15T08:01:00Z", "2026-07-15T08:02:00Z", "2026-07-15T08:03:00Z"]
    client.post("/__admin/drive/reset", headers=ADMIN)
    assert horodatages() == premiers


def test_supprimer_puis_redeposer_donne_un_NOUVEL_element(client, auth):
    avant = element(client, auth, "balance_BOG_2026-01.xlsx")
    assert (
        client.delete(
            "/__admin/drive/files/Insights360/balance_BOG_2026-01.xlsx", headers=ADMIN
        ).status_code
        == 204
    )
    assert client.get(f"/v1.0/drives/{DRIVE}/items/{avant['id']}", headers=auth).status_code == 404
    assert "balance_BOG_2026-01.xlsx" not in {e["name"] for e in tous_les_enfants(client, auth)}
    assert (
        client.delete(
            "/__admin/drive/files/Insights360/balance_BOG_2026-01.xlsx", headers=ADMIN
        ).status_code
        == 404
    )

    contenu = telecharger(client, auth, element(client, auth, "balance_BOG_2026-02.xlsx")["id"])
    assert deposer(client, "Insights360/balance_BOG_2026-01.xlsx", contenu).status_code == 201
    assert element(client, auth, "balance_BOG_2026-01.xlsx")["id"] != avant["id"]


def test_reset_du_lecteur_rend_le_jeu_par_defaut(client, auth):
    avant = {e["name"]: e["cTag"] for e in tous_les_enfants(client, auth)}
    deposer(client, "Insights360/balance_NTE_2026-06.xlsx", b"autre")
    deposer(client, "Insights360/notes.csv", b"x")
    client.delete("/__admin/drive/files/Insights360/taux_2026.xlsx", headers=ADMIN)
    assert client.post("/__admin/drive/reset", headers=ADMIN).status_code == 200
    assert {e["name"]: e["cTag"] for e in tous_les_enfants(client, auth)} == avant


def test_les_noms_que_sharepoint_refuse_sont_refuses(client):
    assert deposer(client, "Insights360/a:b.xlsx", b"x").status_code == 400
    assert deposer(client, "Insights360/a_vti_b.xlsx", b"x").status_code == 400
    assert deposer(client, "Insights360/ espace.xlsx", b"x").status_code == 400
    assert deposer(client, "Insights360/modeles", b"x").status_code == 409


def test_les_compteurs_ne_retiennent_que_les_appels_servis(client, auth):
    """Ce qu'il faut pour PROUVER qu'un second run idempotent ne retélécharge
    rien : les compteurs, pas le contenu des tables."""
    client.post("/__admin/drive/reset", headers=ADMIN)
    client.get(DOSSIER)  # 401 : non compté
    item = client.get(f"{DOSSIER}?$select=id,name", headers=auth).json()["value"][0]
    telecharger(client, auth, item["id"])
    compteurs = client.get("/__admin/drive/counters", headers=ADMIN).json()
    assert compteurs["children"] == 1
    assert compteurs["content"] == 1
    assert compteurs["download"] == 1
    assert compteurs["site"] == 0


def test_le_reset_global_remet_aussi_la_cadence_d_etranglement(client, auth, monkeypatch):
    monkeypatch.setattr(app_module, "THROTTLE_EVERY", 2)
    monkeypatch.setattr(app_module, "_appels", 0)
    assert client.get(SITE, headers=auth).status_code == 200
    assert client.post("/__admin/reset", headers=ADMIN).status_code == 200
    assert client.get(SITE, headers=auth).status_code == 200  # 1er appel après le reset
    assert client.get(SITE, headers=auth).status_code == 429


# ── Le préambule commun : jeton et étranglement ──────────────────────────────

ROUTES_GRAPH = [
    SITE,
    f"/v1.0/sites/{drive_module.SITE_ID}",
    f"/v1.0/sites/{drive_module.SITE_ID}/drive",
    f"/v1.0/sites/{drive_module.SITE_ID}/drives",
    f"/v1.0/drives/{DRIVE}",
    f"/v1.0/drives/{DRIVE}/root",
    f"/v1.0/drives/{DRIVE}/root/children",
    DOSSIER,
    f"/v1.0/drives/{DRIVE}/root:/Insights360/taux_2026.xlsx",
    f"/v1.0/drives/{DRIVE}/root:/Insights360/taux_2026.xlsx:/content",
    f"/v1.0/drives/{DRIVE}/items/01QUELCONQUE",
    f"/v1.0/drives/{DRIVE}/items/01QUELCONQUE/children",
    f"/v1.0/drives/{DRIVE}/items/01QUELCONQUE/content",
]


@pytest.mark.parametrize("url", ROUTES_GRAPH)
def test_le_Bearer_est_exige_sur_chaque_route_du_lecteur(client, url):
    r = client.get(url, follow_redirects=False)
    assert r.status_code == 401
    assert r.json()["error"]["code"] == "InvalidAuthenticationToken"


@pytest.mark.parametrize("url", [SITE, DOSSIER, ROUTES_GRAPH[9]])
def test_l_etranglement_injecte_frappe_aussi_le_lecteur(client, auth, monkeypatch, url):
    monkeypatch.setattr(app_module, "THROTTLE_EVERY", 1)
    monkeypatch.setattr(app_module, "_appels", 0)
    r = client.get(url, headers=auth, follow_redirects=False)
    assert r.status_code == 429
    assert r.headers["Retry-After"] == app_module.RETRY_AFTER
    assert r.json()["error"]["code"] == "TooManyRequests"


# ── Le jeu de données ────────────────────────────────────────────────────────

_MOTIF_BALANCE = re.compile(r"balance_([A-Z]{3})_(\d{4}-\d{2})(?:_[a-z_]+)?\.xlsx")
_MOTIF_TAUX = re.compile(r"taux_\d{4}(?:_[a-z]+)?\.xlsx")
_FORMULE_SOMME = re.compile(r"=(\d+\.\d{2})\+(\d+\.\d{2})")


def _centimes(valeur: Any) -> int | None:
    """Un montant en centimes — None s'il n'est pas un nombre au centime."""
    if isinstance(valeur, str) and (somme := _FORMULE_SOMME.fullmatch(valeur)):
        return int((Decimal(somme[1]) + Decimal(somme[2])) * 100)
    if isinstance(valeur, bool) or not isinstance(valeur, int | float):
        return None
    montant = Decimal(repr(valeur))
    if montant < 0 or montant.as_tuple().exponent < -2:  # type: ignore[operator]
        return None
    return int(montant * 100)


def _defauts_taux(rangees: list[tuple[Any, ...]]) -> set[str]:
    trouves = set() if rangees[0] == jeu.ENTETE_TAUX else {"entete"}
    for type_taux, periode, devise, taux in rangees[1:]:
        if devise == "EUR":
            trouves.add("eur")
        if type_taux not in {"moyen", "budget"} or not isinstance(periode, str):
            trouves.add("periode")
        if not isinstance(taux, float | int) or taux <= 0:
            trouves.add("taux")
    return trouves


def _defauts_lignes(code: str, mois: str, rangees: list[tuple[Any, ...]]) -> set[str]:
    entite = jeu.ENTITES.get(code)
    trouves: set[str] = set() if entite else {"entite"}
    debits = credits = 0
    comptes: list[str] = []
    for code_ligne, mois_ligne, compte, _libelle, debit, credit, devise in rangees:
        if code_ligne != code or mois_ligne != mois:
            trouves.add("entite")
        if entite is not None and devise != entite.devise:
            trouves.add("devise")
        if any(isinstance(v, str) and v.startswith("=") for v in (debit, credit)):
            trouves.add("formule")
        d, c = _centimes(debit), _centimes(credit)
        if not isinstance(compte, str) or d is None or c is None:
            trouves.add("montant")
            continue
        comptes.append(compte)
        debits, credits = debits + d, credits + c
    if debits != credits:
        trouves.add("desequilibre")
    if any(a != b and b.startswith(a) for a in comptes for b in comptes):
        trouves.add("prefixe")
    if len(comptes) != len(set(comptes)):
        trouves.add("doublon")
    return trouves


def defauts(nom: str, contenu: bytes) -> set[str]:
    """Le validateur que le consommateur devra être — en version de référence.

    Rend l'ensemble des règles violées : vide pour un fichier valide.
    """
    classeur = (
        load_workbook(BytesIO(contenu))  # PAS data_only : une formule se voit
        if _MOTIF_BALANCE.fullmatch(nom) or _MOTIF_TAUX.fullmatch(nom)
        else None
    )
    if classeur is None:
        return {"nom"}
    if _MOTIF_TAUX.fullmatch(nom):
        return _defauts_taux(list(classeur["taux"].iter_rows(values_only=True)))
    correspondance = _MOTIF_BALANCE.fullmatch(nom)
    assert correspondance is not None
    rangees = list(classeur["balance"].iter_rows(values_only=True))
    trouves = set() if rangees[0] == jeu.ENTETE_BALANCE else {"entete"}
    return trouves | _defauts_lignes(correspondance[1], correspondance[2], rangees[1:])


def test_chaque_balance_servie_est_valide_equilibree_et_sans_prefixe(client, auth):
    """Les 18 balances et le fichier des taux, tels qu'un consommateur les
    obtient : listés en suivant le curseur, téléchargés par la redirection."""
    vus: set[tuple[str, str]] = set()
    for item in tous_les_enfants(client, auth):
        if "folder" in item:
            continue
        contenu = telecharger(client, auth, item["id"])
        assert defauts(item["name"], contenu) == set(), item["name"]
        if correspondance := _MOTIF_BALANCE.fullmatch(item["name"]):
            vus.add((correspondance[1], correspondance[2]))
    assert vus == {(code, mois) for code in ("NTE", "BOG", "MTL") for mois in jeu.MOIS}


def test_les_balances_ont_les_ordres_de_grandeur_annonces():
    """Recettes au crédit, charges au débit, et des montants plausibles."""
    plages = {"NTE": (130_000, 170_000), "BOG": (360e6, 440e6), "MTL": (105_000, 135_000)}
    for code, (bas, haut) in plages.items():
        for mois in jeu.MOIS:
            lignes = {lg.compte: lg for lg in jeu.balance(code, mois)}
            ca = jeu.ENTITES[code].comptes["ca"][0]
            assert lignes[ca].debit == 0 and bas <= lignes[ca].credit / 100 <= haut
            salaires = jeu.ENTITES[code].comptes["salaires"][0]
            assert lignes[salaires].credit == 0 and lignes[salaires].debit > 0


def test_l_intragroupe_se_repond_a_l_arrondi_de_change_pres():
    """Ce que Bogota et Montréal facturent à Nantes est ce que Nantes achète."""
    for mois in jeu.MOIS:
        nantes = {lg.compte: lg for lg in jeu.balance("NTE", mois)}
        bogota = {lg.compte: lg for lg in jeu.balance("BOG", mois)}
        montreal = {lg.compte: lg for lg in jeu.balance("MTL", mois)}
        taux = jeu.TAUX_MOYENS[mois]
        en_euros = Decimal(bogota["413595"].credit) / Decimal(taux["COP"]) + Decimal(
            montreal["4090"].credit
        ) / Decimal(taux["CAD"])
        assert abs(en_euros - nantes["604800"].debit) <= 1


def test_le_fichier_des_taux(client, auth):
    contenu = telecharger(client, auth, element(client, auth, "taux_2026.xlsx")["id"])
    rangees = list(load_workbook(BytesIO(contenu))["taux"].iter_rows(values_only=True))
    assert rangees[0] == ("type_taux", "periode", "devise", "devise_pour_1_eur")
    moyens = [r for r in rangees[1:] if r[0] == "moyen"]
    assert {(r[1], r[2]) for r in moyens} == {(m, d) for m in jeu.MOIS for d in ("CAD", "COP")}
    assert {(r[1], r[2]) for r in rangees[1:] if r[0] == "budget"} == {
        ("2026", "CAD"),
        ("2026", "COP"),
    }
    for _, _, devise, taux in rangees[1:]:
        assert (4400 <= taux <= 4700) if devise == "COP" else (1.45 <= taux <= 1.52)


def test_le_jeu_est_reproductible_A_L_OCTET():
    """Deux constructions, mêmes octets : l'instantané d'idempotence du
    consommateur (empreintes, tailles) ne doit jamais bouger au redémarrage."""

    def empreintes() -> dict[str, str]:
        rendu = {
            f.chemin: hashlib.sha256(f.contenu).hexdigest() for f in jeu.construire_jeu_par_defaut()
        }
        rendu |= {f.nom: hashlib.sha256(f.fabriquer()).hexdigest() for f in jeu.FIXTURES}
        return rendu

    assert empreintes() == empreintes()


def test_aucun_horodatage_mural_dans_les_classeurs():
    """openpyxl réécrit `dcterms:modified` à chaque save, et zipfile date ses
    membres à l'instant : les deux sont neutralisés."""
    fichier = jeu.jeu_par_defaut()[0]
    with zipfile.ZipFile(BytesIO(fichier.contenu)) as archive:
        core = archive.read("docProps/core.xml").decode()
        assert f">{fichier.modifie}</dcterms:modified>" in core
        assert {info.date_time for info in archive.infolist()} == {(2026, 2, 6, 8, 14, 0)}


# ── Les variantes ────────────────────────────────────────────────────────────

ATTENDUS = {
    "balance_BOG_2026-07_desequilibree.xlsx": {"desequilibre"},
    "balance_MTL_2026-07_formule.xlsx": {"formule"},
    "balance_NTE_2026-07_sous_total.xlsx": {"prefixe"},
    "balance_BOG_2026-07_mauvaise_devise.xlsx": {"devise"},
    "balance_NTE_2026-07_entete.xlsx": {"entete"},
    "balance_XXX_2026-07.xlsx": {"entite"},
    "taux_2026_eur.xlsx": {"eur"},
    "~$balance_NTE_2026-01.xlsx": {"nom"},
    "notes.csv": {"nom"},
    "balance_NTE_2026-07.xlsx": set(),
    "balance_NTE_2026-06_corrigee.xlsx": set(),
}


def test_chaque_variante_porte_SON_defaut_et_un_seul(client):
    """Un test du consommateur qui attend « refusé pour déséquilibre » ne doit
    pas passer parce que le fichier était AUSSI refusé pour son en-tête."""
    index = client.get("/__fixtures/depot_finance").json()["fixtures"]
    assert {f["name"] for f in index} == set(ATTENDUS)
    for f in index:
        r = client.get(_relatif(f["url"]))
        assert r.status_code == 200, f["name"]
        assert defauts(f["name"], r.content) == ATTENDUS[f["name"]], f["name"]
        assert f["valid"] == (not ATTENDUS[f["name"]])


def test_le_fichier_de_verrou_n_est_pas_un_classeur(client):
    contenu = client.get("/__fixtures/depot_finance/~$balance_NTE_2026-01.xlsx").content
    assert len(contenu) == 165
    with pytest.raises(zipfile.BadZipFile):
        zipfile.ZipFile(BytesIO(contenu))


def test_la_formule_n_a_pas_de_valeur_en_cache(client):
    """Un lecteur en `data_only=True` lit None — et doit refuser, lui aussi."""
    contenu = client.get("/__fixtures/depot_finance/balance_MTL_2026-07_formule.xlsx").content
    rangees = list(
        load_workbook(BytesIO(contenu), data_only=True)["balance"].iter_rows(values_only=True)
    )
    assert any(r[4] is None for r in rangees[1:])


def test_aucune_variante_invalide_dans_le_dossier_par_defaut(client, auth):
    """Elles rendraient rouge la CI du consommateur : on les dépose à la main."""
    noms = {e["name"] for e in tous_les_enfants(client, auth)}
    assert not noms & set(ATTENDUS)
    assert client.get("/__fixtures/depot_finance/inconnue.xlsx").status_code == 404


# ── Le contrat ───────────────────────────────────────────────────────────────


def test_le_contrat_publie_est_a_jour_et_couvre_le_lecteur():
    """`make contract` doit avoir été relancé : le consommateur en épingle une
    copie, et un contrat périmé lui ferait valider contre une surface fausse."""
    from entra_mock import app

    publie = yaml.safe_load(Path("contracts/msgraph.openapi.yaml").read_text())
    assert publie == app.openapi()
    chemins = set(publie["paths"])
    assert {
        "/v1.0/sites/{hostname}:/{server_relative_path}",
        "/v1.0/sites/{site_id}/drive",
        "/v1.0/sites/{site_id}/drives",
        "/v1.0/drives/{drive_id}/root:/{chemin}",
        "/v1.0/drives/{drive_id}/items/{item_id}",
        "/v1.0/drives/{drive_id}/items/{item_id}/content",
    } <= chemins
    assert not any(c.startswith(("/__admin", "/__fixtures", "/sites/")) for c in chemins)
    assert (
        "302"
        in publie["paths"]["/v1.0/drives/{drive_id}/items/{item_id}/content"]["get"]["responses"]
    )
